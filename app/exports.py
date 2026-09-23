"""Themed protocol exports based on the supplied secretary's Word template."""
from __future__ import annotations
from copy import deepcopy
from html import escape
from io import BytesIO
import os
from pathlib import Path
import re
from typing import Any

TEMPLATE = Path(__file__).with_name("assets") / "protocol-template.docx"
LABELS = {"action": "Поручение", "decision": "Решение", "initiative": "Инициатива",
          "question": "Открытый вопрос", "risk": "Риск"}

def _time(seconds: float) -> str:
    return f"{int(seconds // 60):02d}:{int(seconds % 60):02d}"

def sections(meeting: dict[str, Any]) -> tuple[list[dict], list[dict]]:
    """Associate tasks through cited speech, never through hardcoded agenda labels."""
    topics = meeting.get("summary_topics") or []
    if not topics:
        for block in (meeting.get("summary") or "").split("\n\n"):
            heading, separator, text = block.strip().partition("\n")
            if heading:
                topics.append({"title": heading if separator else "Итоги обсуждения",
                               "text": text if separator else heading, "source_segment_ids": []})
    parts = [{"title": re.sub(r"^(?:часть|тема)\s*\d+\s*[.:—–-]?\s*", "", str(t["title"]), flags=re.I),
              "text": t["text"], "source_segment_ids": list(map(str, t.get("source_segment_ids", []))),
              "items": []} for t in topics]
    unassigned = []
    for item in meeting["items"]:
        ids = set(map(str, item.get("source_segment_ids", [])))
        scores = [len(ids.intersection(t["source_segment_ids"])) for t in parts]
        if scores and max(scores):
            parts[scores.index(max(scores))]["items"].append(item)
        elif len(parts) == 1:
            parts[0]["items"].append(item)
        else:
            unassigned.append(item)
    return parts, unassigned

def _task_rows(items):
    return [[("Отменено: " if i.get("status") == "cancelled" else "") + i["title"],
             i.get("owner") or "Требует уточнения", i.get("due_text") or "Не указан"]
            for i in items if i["kind"] == "action"]

def _items_flow(items):
    rows = _task_rows(items)
    if rows:
        yield "section", "Поручения"
        yield "table", rows
    for item in items:
        if item["kind"] != "action":
            yield "body", f"{LABELS.get(item['kind'], item['kind'])}: {item['title']}"

def _transcript_flow(meeting, parts):
    """Keep every turn once, in recording order, under actual agenda headings."""
    active = None
    seen = set()
    for segment in sorted(meeting["segments"], key=lambda s: s["start"]):
        matches = [i for i, part in enumerate(parts) if str(segment["id"]) in part["source_segment_ids"]]
        topic = active if active in matches else matches[0] if matches else active
        if topic != active or not seen:
            if topic is None:
                yield "heading", "Текст совещания"
            else:
                continued = " · продолжение" if topic in seen else ""
                yield "heading", f"Часть {topic + 1}. {parts[topic]['title']}{continued}"
            seen.add(topic)
            active = topic
        speaker = meeting["speakers"].get(segment["speaker_id"], segment["speaker_id"])
        yield "speaker", f"{speaker} · {_time(segment['start'])}"
        yield "body", segment["text"]

def _flow(meeting, include_transcript):
    yield "title", "Протокол совещания"
    yield "organization", os.getenv("PROTOCOL_ORGANIZATION", "АО «Самрук-Қазына»")
    yield "subject", "Тема: " + meeting["title"]
    status = "Утверждён" if meeting["state"] == "approved" else "ЧЕРНОВИК · не утверждён"
    yield "metadata", f"Дата совещания: {meeting['meeting_date']} · {status}"
    parts, remaining = sections(meeting)
    if include_transcript:
        yield from _transcript_flow(meeting, parts)
        yield "pagebreak", ""
    yield "heading", "Саммари по ключевым пунктам"
    for index, part in enumerate(parts, 1):
        yield "section", f"Тема {index}. {part['title']}"
        yield "body", part["text"]
        yield from _items_flow(part["items"])
    if remaining:
        yield "section", "Общие поручения и решения"
        yield from _items_flow(remaining)

def docx(meeting: dict[str, Any], *, include_transcript: bool = True) -> bytes:
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt
    from docx.text.paragraph import Paragraph
    from docx.table import Table
    document = Document(TEMPLATE)
    patterns = [deepcopy(p._p) for p in document.paragraphs]
    table_pattern = deepcopy(document.tables[0]._tbl)
    body = document._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)
    roles = {"title": 0, "organization": 1, "subject": 2, "metadata": 3,
             "heading": 4, "section": 5, "body": 6, "speaker": 6}
    for kind, value in _flow(meeting, include_transcript):
        if kind == "pagebreak":
            document.add_page_break()
        elif kind == "table":
            element = deepcopy(table_pattern)
            body.insert(len(body) - 1, element)
            table = Table(element, document._body)
            row_pattern = deepcopy(table.rows[1]._tr)
            element.remove(table.rows[1]._tr)
            table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
            for values in value:
                element.append(deepcopy(row_pattern))
                row = table.rows[-1]
                for cell, text in zip(row.cells, values):
                    cell.text = text
                row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
            for row_index, row in enumerate(table.rows):
                for cell in row.cells:
                    for p in cell.paragraphs:
                        p.paragraph_format.space_after = Pt(4)
                        p.paragraph_format.space_before = Pt(4)
                        for run in p.runs:
                            run.font.name = "Times New Roman"
                            run.font.size = Pt(11)
                            run.bold = row_index == 0
        else:
            index = roles[kind]
            element = deepcopy(patterns[index])
            body.insert(len(body) - 1, element)
            p = Paragraph(element, document._body)
            for run in p.runs:
                run.text = ""
            run = p.runs[0] if p.runs else p.add_run()
            run.text = value
            if index in {1, 2, 3, 6}:
                run.font.name = "Times New Roman"
                run.font.size = Pt(9 if index == 3 else 11)
            if kind == "speaker":
                run.bold = True
                p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.keep_with_next = kind != "body"
            p.paragraph_format.widow_control = True
    output = BytesIO()
    document.save(output)
    return output.getvalue()
def _font_path() -> Path:
    candidates = [os.getenv("PROTOCOL_PDF_FONT", ""),
                  "/System/Library/Fonts/Supplemental/Arial Unicode.ttf",
                  "/Library/Fonts/Arial Unicode.ttf",
                  "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
                  "/usr/share/fonts/truetype/noto/NotoSans-Regular.ttf"]
    for candidate in candidates:
        if candidate and Path(candidate).exists():
            return Path(candidate)
    raise RuntimeError("Не найден Unicode-шрифт для PDF. Задайте PROTOCOL_PDF_FONT")

def pdf(meeting: dict[str, Any], *, include_transcript: bool = True) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak
    pdfmetrics.registerFont(TTFont("ProtocolUnicode", str(_font_path())))
    normal = ParagraphStyle("body", fontName="ProtocolUnicode", fontSize=10.5, leading=14, spaceAfter=7)
    title = ParagraphStyle("title", parent=normal, fontSize=25, leading=30, alignment=TA_CENTER,
                           spaceAfter=8, keepWithNext=True)
    center = ParagraphStyle("center", parent=normal, alignment=TA_CENTER, keepWithNext=True)
    metadata = ParagraphStyle("metadata", parent=center, fontSize=8, leading=11)
    heading = ParagraphStyle("heading", parent=normal, fontSize=15, leading=19, spaceBefore=15,
                             spaceAfter=9, textColor=colors.HexColor("#2E74B5"), keepWithNext=True)
    subheading = ParagraphStyle("part", parent=heading, fontSize=12.5, leading=16, spaceBefore=13)
    speaker = ParagraphStyle("speaker", parent=normal, spaceBefore=7, spaceAfter=3, keepWithNext=True)
    cell = ParagraphStyle("cell", parent=normal, fontSize=10, leading=13, spaceAfter=0)
    styles = {"title": title, "organization": center, "subject": center, "metadata": metadata,
              "heading": heading, "section": subheading, "body": normal, "speaker": speaker}
    def p(value, style=normal):
        return Paragraph(escape(str(value)).replace("\n", "<br/>"), style)
    story = []
    for kind, value in _flow(meeting, include_transcript):
        if kind == "pagebreak":
            story.append(PageBreak())
        elif kind == "table":
            rows = [[p(v, cell) for v in ["Поручение", "Ответственный", "Срок"]]]
            rows.extend([[p(v, cell) for v in row] for row in value])
            table = Table(rows, colWidths=[210.6, 140.4, 117], repeatRows=1, hAlign="LEFT")
            table.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#D9E2F3")),
                ("GRID", (0, 0), (-1, -1), .4, colors.HexColor("#657383")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 6), ("RIGHTPADDING", (0, 0), (-1, -1), 6),
                ("TOPPADDING", (0, 0), (-1, -1), 6), ("BOTTOMPADDING", (0, 0), (-1, -1), 6)]))
            story.extend([table, Spacer(1, 8)])
        else:
            story.append(p(value, styles[kind]))
    output = BytesIO()
    SimpleDocTemplate(output, pagesize=letter, leftMargin=72, rightMargin=72,
                      topMargin=72, bottomMargin=72, title="Протокол совещания").build(story)
    return output.getvalue()
